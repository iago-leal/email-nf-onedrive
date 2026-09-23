"""Sumário LEIAME da raiz do destino (feature 006).

Os documentos sem vencimento identificado ficam na raiz do destino. Para que o financeiro saiba
o que são e por que não foram classificados, a raiz ganha a planilha `NOME_LEIAME`, cujo "00"
a põe no topo da listagem. Ela é montada da listagem real da raiz, cruzada com o registro,
e refeita sempre que o conjunto de nomes da raiz muda, inclusive quando alguém move ou apaga
um arquivo à mão; é o único arquivo que a ferramenta sobrescreve (`Rclone.publicar_leiame`).
"""

from __future__ import annotations

import hashlib
import logging
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from email_nf_onedrive.coleta.coleta import Falha
from email_nf_onedrive.envio.rclone import NOME_LEIAME, ErroRclone, Rclone
from email_nf_onedrive.envio.vencimento import PASTA_VENCIMENTOS
from email_nf_onedrive.registro.banco import Registro

BRASILIA = timezone(timedelta(hours=-3), "Brasília")  # sem horário de verão desde 2019
META_ASSINATURA = "leiame:"
MOTIVO_FORA_DA_FERRAMENTA = "não enviado por esta ferramenta (posto à mão na pasta)"
MOTIVO_NAO_APURADO = "motivo não apurado (enviado antes do sumário existir)"
COLUNAS = (("Arquivo", 60), ("Empresa", 10), ("Fornecedor", 32), ("Motivo", 60),
           ("Recebido em", 17), ("Remetente", 34), ("Assunto", 50))
LINHA_CABECALHO = 5


@dataclass(frozen=True)
class Linha:
    arquivo: str
    empresa: str
    fornecedor: str
    motivo: str
    recebido_em: datetime | None
    remetente: str
    assunto: str


def _partes_do_nome(nome: str) -> tuple[str, str]:
    """Empresa e fornecedor pelo padrão `EMPRESA - FORNECEDOR [NF n] - TIPO.ext`; vazios fora dele."""
    partes = Path(nome).stem.split(" - ")
    if len(partes) < 3:
        return "", ""
    fornecedor = partes[1].split(" NF ")[0]
    return partes[0], fornecedor


def linhas_da_raiz(arquivos: list[str], registro: Registro, destino: str) -> list[Linha]:
    enviados = registro.enviados_na_pasta(destino)
    linhas = []
    for nome in arquivos:
        empresa, fornecedor = _partes_do_nome(nome)
        if nome in enviados:
            anexo, motivo = enviados[nome]
            linhas.append(Linha(nome, empresa, fornecedor, motivo or MOTIVO_NAO_APURADO, anexo.data_mensagem,
                                anexo.remetente, anexo.assunto))
        else:
            linhas.append(Linha(nome, empresa, fornecedor, MOTIVO_FORA_DA_FERRAMENTA, None, "", ""))
    return sorted(linhas, key=lambda l: (l.motivo, l.empresa, l.arquivo))


def montar_planilha(linhas: list[Linha], agora: datetime, saida: Path) -> None:
    livro = Workbook()
    folha = livro.active
    folha.title = "Documentos"
    folha["A1"] = "LEIAME: documentos sem vencimento identificado"
    folha["A1"].font = Font(bold=True, size=14)
    folha["A2"] = ("Estes documentos ficaram fora da grade porque a ferramenta não identificou o vencimento "
                   f"deles. Para classificar um documento, mova-o para {PASTA_VENCIMENTOS}/DIA <dia>/202X-<mês>. "
                   "Esta planilha é refeita automaticamente a cada mudança nesta pasta; alterações feitas "
                   "nela se perdem.")
    folha["A3"] = (f"Atualizada em {agora.astimezone(BRASILIA):%d/%m/%Y %H:%M} (horário de Brasília). "
                   f"{len(linhas)} documento(s) nesta pasta.")
    negrito = Font(bold=True)
    fundo = PatternFill("solid", fgColor="DDEBF7")
    for coluna, (titulo, largura) in enumerate(COLUNAS, start=1):
        celula = folha.cell(LINHA_CABECALHO, coluna, titulo)
        celula.font, celula.fill = negrito, fundo
        folha.column_dimensions[get_column_letter(coluna)].width = largura
    for n, linha in enumerate(linhas, start=LINHA_CABECALHO + 1):
        recebido = linha.recebido_em.astimezone(BRASILIA).replace(tzinfo=None) if linha.recebido_em else None
        valores = (linha.arquivo, linha.empresa, linha.fornecedor, linha.motivo, recebido, linha.remetente,
                   linha.assunto)
        for coluna, valor in enumerate(valores, start=1):
            folha.cell(n, coluna, valor)
        folha.cell(n, 5).number_format = "dd/mm/yyyy hh:mm"
    ultima = LINHA_CABECALHO + max(len(linhas), 1)
    folha.auto_filter.ref = f"A{LINHA_CABECALHO}:{get_column_letter(len(COLUNAS))}{ultima}"
    folha.freeze_panes = f"A{LINHA_CABECALHO + 1}"

    resumo = livro.create_sheet("Resumo por motivo")
    resumo.append(("Motivo", "Documentos"))
    for celula in resumo[1]:
        celula.font, celula.fill = negrito, fundo
    for motivo, quantidade in sorted(Counter(l.motivo for l in linhas).items(), key=lambda par: -par[1]):
        resumo.append((motivo, quantidade))
    resumo.column_dimensions["A"].width = 80
    resumo.column_dimensions["B"].width = 12
    for celula in resumo["A"]:
        celula.alignment = Alignment(wrap_text=True)
    livro.save(saida)


def _assinatura(arquivos: list[str]) -> str:
    return hashlib.sha256("\n".join(arquivos).encode()).hexdigest()


def atualizar_leiame(destino: str, registro: Registro, rclone: Rclone, trabalho: Path, logger: logging.Logger,
                     *, agora: datetime, simulacao: bool = False) -> Falha | None:
    """Refaz o sumário da raiz de `destino` se a raiz mudou desde o último; devolve a falha, se houver."""
    destino = destino.rstrip("/")
    try:
        nomes = rclone.listar(destino)
        if nomes is None:
            return None  # o envio já avisa do destino ausente
        arquivos = sorted(n for n in nomes if not n.endswith("/") and n != NOME_LEIAME)
        assinatura = _assinatura(arquivos)
        if NOME_LEIAME in nomes and registro.meta_obter(META_ASSINATURA + destino) == assinatura:
            return None
        if simulacao:
            logger.info("simulação: refaria o sumário %s com %d documento(s)", NOME_LEIAME, len(arquivos))
            return None
        trabalho.mkdir(parents=True, exist_ok=True)
        local = trabalho / NOME_LEIAME
        montar_planilha(linhas_da_raiz(arquivos, registro, destino), agora, local)
        rclone.publicar_leiame(local, f"{destino}/{NOME_LEIAME}")
        local.unlink(missing_ok=True)
    except ErroRclone as erro:
        logger.error("sumário %s não atualizado: %s: %s", NOME_LEIAME, erro.causa, erro.detalhe)
        return Falha("onedrive:leiame", f"OneDrive: o sumário {NOME_LEIAME} de {destino} não foi atualizado "
                                        f"({erro.causa}); a próxima execução tenta de novo.")
    registro.meta_definir(META_ASSINATURA + destino, assinatura)
    registro.commit()
    logger.info("sumário %s refeito: %d documento(s) na raiz", NOME_LEIAME, len(arquivos))
    return None
