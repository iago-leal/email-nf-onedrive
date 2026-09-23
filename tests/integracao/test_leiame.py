"""Sumário LEIAME da raiz do destino, com o Rclone real contra diretório local (adendo 006)."""

from __future__ import annotations

import logging
import shutil
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from email_nf_onedrive.envio.envio import enviar_anexos
from email_nf_onedrive.envio.leiame import (
    LINHA_CABECALHO, MOTIVO_FORA_DA_FERRAMENTA, MOTIVO_NAO_APURADO, atualizar_leiame,
)
from email_nf_onedrive.envio.rclone import NOME_LEIAME, ErroRclone, Rclone
from email_nf_onedrive.envio.vencimento import MOTIVO_SEM_FATURA, MOTIVO_SERVICO
from tests.conftest import nfe_com_vencimentos, pdf_com_texto

from .test_envio import destino, novo_item, registro  # noqa: F401 (fixtures)

pytestmark = [
    pytest.mark.integracao,
    pytest.mark.skipif(shutil.which("rclone") is None, reason="rclone não instalado"),
]

LOG = logging.getLogger("teste-leiame")
AGORA = datetime(2026, 9, 23, 18, 30, tzinfo=timezone.utc)
NFSE = pdf_com_texto("NFS-e Nota Fiscal de Servico Eletronica",
                     "Prefeitura Municipal emitiu este documento para o tomador indicado conforme contrato vigente")


def _atualizar(registro, destino, trabalho, rclone=None, **opcoes):
    return atualizar_leiame(str(destino), registro, rclone or Rclone(":local"), trabalho, LOG, agora=AGORA, **opcoes)


def _linhas(destino) -> list[tuple]:
    folha = load_workbook(destino / NOME_LEIAME)["Documentos"]
    return [linha for linha in folha.iter_rows(min_row=LINHA_CABECALHO + 1, values_only=True)]


def test_sumario_lista_a_raiz_com_o_motivo_de_cada_documento(registro, destino, novo_item, tmp_path):
    enviar_anexos([novo_item(NFSE, nome="nfse.pdf", remetente="nfse@prefeitura.example"),
                   novo_item(nfe_com_vencimentos(), nome="nota.xml", classe="nfe-xml"),
                   novo_item(nfe_com_vencimentos("2026-10-05"), nome="outra.xml", classe="nfe-xml")],
                  registro, Rclone(":local"), LOG)
    (destino / "posto a mao.pdf").write_bytes(b"%PDF manual")
    assert _atualizar(registro, destino, tmp_path / "trabalho") is None

    linhas = _linhas(destino)
    assert [(l[0], l[1], l[2], l[3]) for l in linhas] == [
        ("ACME - PREFEITURA - BOLETO.pdf", "ACME", "PREFEITURA", MOTIVO_SERVICO),
        ("ACME - FORNECEDOR FICTICIO LTDA NF 1234 - REF.xml", "ACME", "FORNECEDOR FICTICIO LTDA", MOTIVO_SEM_FATURA),
        ("posto a mao.pdf", None, None, MOTIVO_FORA_DA_FERRAMENTA),   # célula vazia volta como None
    ]
    assert linhas[0][4] == datetime(2026, 9, 18, 12, 0)   # 15:00 UTC em Brasília, sem fuso na célula
    assert linhas[0][5:] == ("nfse@prefeitura.example", "Boleto")
    assert linhas[2][4:] == (None, None, None)
    livro = load_workbook(destino / NOME_LEIAME)
    assert "23/09/2026 15:30" in livro["Documentos"]["A3"].value
    resumo = list(livro["Resumo por motivo"].iter_rows(min_row=2, values_only=True))
    assert sorted(resumo) == sorted([(MOTIVO_SEM_FATURA, 1), (MOTIVO_SERVICO, 1), (MOTIVO_FORA_DA_FERRAMENTA, 1)])
    assert list((tmp_path / "trabalho").iterdir()) == []


class _RcloneContado(Rclone):
    def __init__(self) -> None:
        super().__init__(":local")
        self.publicacoes = 0

    def publicar_leiame(self, local, relativo) -> None:
        self.publicacoes += 1
        super().publicar_leiame(local, relativo)


def test_sumario_so_e_refeito_quando_a_raiz_muda_ou_ele_some(registro, destino, tmp_path):
    rclone = _RcloneContado()
    (destino / "a.pdf").write_bytes(b"%PDF a")
    (destino / "PASTA").mkdir()
    _atualizar(registro, destino, tmp_path, rclone)
    _atualizar(registro, destino, tmp_path, rclone)
    assert rclone.publicacoes == 1

    (destino / "b.pdf").write_bytes(b"%PDF b")                     # chegou um documento
    _atualizar(registro, destino, tmp_path, rclone)
    (destino / "a.pdf").rename(destino / "PASTA" / "a.pdf")         # o financeiro classificou um à mão
    _atualizar(registro, destino, tmp_path, rclone)
    assert rclone.publicacoes == 3
    assert [l[0] for l in _linhas(destino)] == ["b.pdf"]

    (destino / NOME_LEIAME).unlink()                                # alguém apagou o sumário
    _atualizar(registro, destino, tmp_path, rclone)
    assert rclone.publicacoes == 4 and (destino / NOME_LEIAME).exists()


def test_documento_enviado_antes_da_feature_aparece_sem_motivo_apurado(registro, destino, novo_item, tmp_path):
    item = novo_item(b"%PDF antigo", nome="antigo.pdf")
    registro.marcar_enviado(item.anexo_id, f"{destino}/ACME - FORNECEDOR - BOLETO.pdf")
    (destino / "ACME - FORNECEDOR - BOLETO.pdf").write_bytes(b"%PDF antigo")
    _atualizar(registro, destino, tmp_path)
    assert _linhas(destino)[0][3] == MOTIVO_NAO_APURADO


def test_simulacao_nao_escreve_o_sumario(registro, destino, tmp_path):
    (destino / "a.pdf").write_bytes(b"%PDF a")
    _atualizar(registro, destino, tmp_path, simulacao=True)
    assert not (destino / NOME_LEIAME).exists()


def test_falha_no_onedrive_vira_aviso_e_tenta_de_novo(registro, destino, tmp_path):
    class _RcloneRecusado(Rclone):
        def publicar_leiame(self, local, relativo) -> None:
            raise ErroRclone("acesso", "403")

    (destino / "a.pdf").write_bytes(b"%PDF a")
    falha = _atualizar(registro, destino, tmp_path, _RcloneRecusado(":local"))
    assert falha.causa == "onedrive:leiame"
    _atualizar(registro, destino, tmp_path)                         # a assinatura não foi gravada
    assert (destino / NOME_LEIAME).exists()
